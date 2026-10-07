"""Who may see, download and change the interactions, and the files that start an empty database."""
import csv
import io
import json
from pathlib import Path

from django.conf import settings
from django.core.management import call_command
from django.urls import reverse

from beetlesgallery.beetles_app import interaction_data as data
from beetlesgallery.beetles_app.interaction_upload import import_interactions
from beetlesgallery.beetles_app.management.commands import make_interactions_upload_file as maker
from beetlesgallery.beetles_app.models import PathogenInteraction
from beetlesgallery.beetles_app.testing import PageBehaviourCase, make_taxon

NAMES = ["dataset", "validated", "fungi", "nematodes", "microsporidia", "protists", "bacteria", "viruses", "hosts", "references"]


def download_url(name, ext="csv"):
    return reverse("interactions_download", args=[name, ext])


class AccessTests(PageBehaviourCase):
    def setUp(self):
        super().setUp()
        PathogenInteraction.objects.create(
            beetle_host="Ips typographus", beetle_host_id="1733", pathogen="Beauveria bassiana", category="Fungi",
            record_number="1", validation_type="virulence test", title="A paper", year="2013")

    def test_viewing_is_free(self):
        self.assertEqual(self.client.get(reverse("interactions_preview")).status_code, 200)
        for name in ("interactions_records", "interactions_hosts", "interactions_references"):
            self.assertEqual(self.client.get(reverse(name)).status_code, 200)

    def test_downloading_needs_an_account(self):
        for name in NAMES:
            self.assertRedirectsToLogin(self.client.get(download_url(name)))
        self.assertRedirectsToLogin(self.client.get(download_url("workbook", "xlsx")))
        self.assertRedirectsToLogin(self.client.get(reverse("interactions_export")))

    def test_any_account_can_download(self):
        self.client.force_login(self.user)   # a plain member
        for name in NAMES:
            response = self.client.get(download_url(name))
            self.assertEqual(response.status_code, 200, name)
            self.assertIn("text/csv", response["Content-Type"])
            self.assertIn(f"beetle_interactions_{name}_", response["Content-Disposition"])

    def test_the_csvs_have_the_published_columns_and_the_rows_of_their_group(self):
        PathogenInteraction.objects.create(beetle_host="Ips typographus", beetle_host_id="1733", pathogen="Nem", category="Nematode", record_number="2")
        self.client.force_login(self.user)

        def rows(name):
            text = self.client.get(download_url(name)).content.decode("utf-8-sig")
            return list(csv.DictReader(io.StringIO(text)))

        self.assertEqual(list(rows("dataset")[0]), data_columns())
        self.assertEqual([r["pathogens"] for r in rows("dataset")], ["Beauveria bassiana", "Nem"])
        self.assertEqual([r["pathogens"] for r in rows("fungi")], ["Beauveria bassiana"])
        self.assertEqual([r["pathogens"] for r in rows("nematodes")], ["Nem"])
        self.assertEqual([r["pathogens"] for r in rows("validated")], ["Beauveria bassiana"])
        self.assertEqual(rows("hosts")[0]["No. records"], "2")
        self.assertEqual(rows("references")[0]["Citation"], "A paper")

    def test_the_excel_workbook_has_three_sheets(self):
        import openpyxl
        self.client.force_login(self.user)
        response = self.client.get(download_url("workbook", "xlsx"))
        self.assertEqual(response.status_code, 200)
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.sheetnames, ["Dataset", "Beetle Hosts", "References"])
        self.assertEqual(book["Dataset"].max_row, 2)

    def test_unknown_downloads_are_not_found(self):
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(download_url("nope")).status_code, 404)
        self.assertEqual(self.client.get(download_url("dataset", "xlsx")).status_code, 404)

    def test_the_page_offers_a_download_to_accounts_and_a_sign_in_to_everyone_else(self):
        anonymous = self.client.get(reverse("interactions_preview")).content.decode()
        self.assertIn("Download CSV (sign in)", anonymous)
        self.assertNotIn(download_url("hosts"), anonymous)
        self.client.force_login(self.user)
        member = self.client.get(reverse("interactions_preview")).content.decode()
        self.assertIn(download_url("hosts"), member)
        self.assertIn(download_url("workbook", "xlsx"), member)
        self.assertIn('id="download-active-csv-btn"', member)
        self.assertNotIn("(sign in)", member)

    def test_the_interaction_pages_go_back_to_the_interactions_page(self):
        self.client.force_login(self.superuser if hasattr(self, "superuser") else self.staff)
        back = f'href="{reverse("interactions_preview")}" class="inline-block py-2'
        for name in ("upload_interactions", "interaction_review", "upload_interaction_proposals"):
            response = self.client.get(reverse(name))
            if response.status_code != 200:
                continue   # the proposals upload is for superusers; checked below
            page = response.content.decode()
            self.assertIn(back, page, name)
            self.assertNotIn("&larr; Data Management", page, name)
        from django.conf import settings
        for template in ("interaction_review", "upload_interaction_proposals", "upload_interactions"):
            source = (settings.BASE_DIR / "beetlesgallery" / "templates" / "beetles" / f"{template}.html").read_text()
            self.assertIn("{% url 'interactions_preview' as back_url %}", source, template)   # site-back, #618
            self.assertNotIn("&larr; Data Management", source, template)

    def test_changing_the_data_needs_staff(self):
        for name in ("upload_interactions", "interactions_initial_file", "interaction_review"):
            self.assertRedirectsToLogin(self.client.get(reverse(name)))
        self.client.force_login(self.user)
        for name in ("upload_interactions", "interactions_initial_file", "interaction_review"):
            self.assertRedirectsToLogin(self.client.get(reverse(name)))
        self.assertRedirectsToLogin(self.client.post(reverse("upload_interactions"), {}))
        self.client.force_login(self.staff)
        for name in ("upload_interactions", "interactions_initial_file", "interaction_review"):
            self.assertEqual(self.client.get(reverse(name)).status_code, 200, name)

    def test_the_published_files_are_not_served_to_the_public(self):
        static = Path(settings.BASE_DIR) / "beetlesgallery" / "static"
        self.assertFalse((static / "data").exists(), "the published dataset must not be under static/")
        self.assertEqual(self.client.get("/static/data/bark_beetle_pathogens_master.json").status_code, 404)


def data_columns():
    from beetlesgallery.beetles_app.interaction_downloads import RECORD_COLUMNS
    return RECORD_COLUMNS


class InitialFileTests(PageBehaviourCase):
    """The file to upload to start an empty database with what is published today."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.published = json.loads(maker.SOURCE.read_text(encoding="utf-8"))

    def give_the_species_list(self):
        seen = {}
        for r in self.published:
            seen.setdefault(r["Beetle Host IDs"], r["Beetle Host"])
        for valid_id, name in seen.items():
            genus, _, species = name.partition(" ")
            make_taxon(valid_species_id=valid_id, scientific_name=name, genus=genus, species=species)

    def test_the_initial_file_download_is_built_from_the_published_json(self):
        self.client.force_login(self.staff)
        response = self.client.get(reverse("interactions_initial_file"))
        self.assertEqual(response.content.decode("utf-8"), maker.render(self.published))

    def test_the_file_has_every_record_and_the_upload_columns(self):
        rows = list(csv.DictReader(io.StringIO(maker.render(self.published).lstrip("\ufeff"))))
        self.assertEqual(len(rows), 1015)
        self.assertEqual(list(rows[0]), maker.COLUMNS)
        self.assertEqual({r["record_id"] for r in rows}, {"NEW"})
        self.assertEqual(sorted(int(r["records_id"]) for r in rows), list(range(1, 1016)))

    def test_uploading_it_into_an_empty_database_recreates_the_published_dataset(self):
        self.give_the_species_list()
        text = maker.render(self.published).lstrip("\ufeff")
        check = import_interactions(text, user=self.staff, dry_run=True)
        self.assertTrue(check.ok, check.errors)
        self.assertEqual(PathogenInteraction.objects.count(), 0)          # a check saves nothing
        result = import_interactions(text, user=self.staff)
        self.assertTrue(result.ok, result.errors)
        self.assertEqual((result.rows, result.created), (1015, 1015))
        stats = data.page_stats()
        self.assertEqual((stats["records"], stats["hosts"], stats["genera"], stats["taxa"], stats["sources"], stats["validated"]),
                         (1015, 108, 28, 274, 281, 245))
        back = data.master_records()
        for original, mine in zip(self.published, back):
            for column, value in original.items():
                if column == "Beetle Host IDs":
                    value = value or ""
                if isinstance(value, str):
                    value = value.strip()   # an upload reads cells trimmed ("non full-text " in the paper's file)
                self.assertEqual(mine[column] or None, value or None, (original["Records ID"], column))

    def test_the_deploys_loader_then_adds_nothing(self):
        self.give_the_species_list()
        import_interactions(maker.render(self.published).lstrip("\ufeff"), user=self.staff)
        call_command("import_pathogen_interactions", stdout=io.StringIO())
        self.assertEqual(PathogenInteraction.objects.count(), 1015)

    def test_uploading_it_twice_is_refused_and_changes_nothing(self):
        self.give_the_species_list()
        text = maker.render(self.published).lstrip("\ufeff")
        import_interactions(text, user=self.staff)
        again = import_interactions(text, user=self.staff)
        self.assertFalse(again.ok)
        self.assertIn("already used", again.errors[0])
        self.assertEqual(PathogenInteraction.objects.count(), 1015)

    def test_the_template_is_a_valid_upload_with_the_same_columns(self):
        template = Path(settings.BASE_DIR) / "beetlesgallery" / "static" / "downloads" / "interactions_upload_template.csv"
        self.assertEqual(next(csv.reader(io.StringIO(template.read_text(encoding="utf-8")))), maker.COLUMNS)
        make_taxon(valid_species_id="1733", scientific_name="Ips typographus", genus="Ips", species="typographus")
        result = import_interactions(template.read_text(encoding="utf-8"), user=self.staff, dry_run=True)
        self.assertTrue(result.ok, result.errors)


class RecordsIdTests(PageBehaviourCase):
    def setUp(self):
        super().setUp()
        make_taxon(valid_species_id="1733", scientific_name="Ips typographus", genus="Ips", species="typographus")

    def load(self, *rows):
        header = "record_id,records_id,beetle_host,pathogen,category\n"
        return import_interactions(header + "\n".join(rows) + "\n", user=self.staff)

    def test_a_number_is_kept_on_a_new_row(self):
        self.assertTrue(self.load("NEW,7,Ips typographus,Beauveria,Fungi").ok)
        self.assertEqual(PathogenInteraction.objects.get().record_number, "7")

    def test_a_used_repeated_or_malformed_number_is_refused(self):
        self.load("NEW,7,Ips typographus,Beauveria,Fungi")
        self.assertIn("already used", self.load("NEW,7,Ips typographus,Other,Fungi").errors[0])
        self.assertIn("repeats row 2", self.load("NEW,8,Ips typographus,A,Fungi", "NEW,8,Ips typographus,B,Fungi").errors[0])
        self.assertIn("whole number", self.load("NEW,x1,Ips typographus,A,Fungi").errors[0])
        self.assertEqual(PathogenInteraction.objects.count(), 1)

    def test_a_number_cannot_be_changed_but_the_same_one_may_stay(self):
        self.load("NEW,7,Ips typographus,Beauveria,Fungi")
        row = PathogenInteraction.objects.get()
        changed = import_interactions(f"record_id,records_id\n{row.id},9\n", user=self.staff)
        self.assertIn("cannot be changed", changed.errors[0])
        ok = import_interactions(f"record_id,records_id,year\n{row.id},7,2020\n", user=self.staff)
        self.assertTrue(ok.ok, ok.errors)
